"""Magello-presentationer i webbens formspråk (spår B i AGENTS.md).

Bygger en .pptx från en tom 16:9-presentation med webbens styleguide som förlaga:
tokens/tokens.json, styleguide/css/magello.css, assets/. Ingen PowerPoint-mall.
Regler och bildtyper beskrivs i docs/presentation-webb.md.

Kräver: python-pptx, Pillow, cairosvg.

    from magello_pptx import *
    s = page()                      # vit bild
    text(s, L, 1000000, W, 300000, [("kicker", "Etikett")])
    ...
    save("presentation.pptx", title="Titel")
"""
import os, sys, tempfile
from PIL import Image, ImageEnhance
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
ASSETS = os.path.join(REPO, "assets")
CACHE = os.path.join(tempfile.gettempdir(), "magello_pptx")
os.makedirs(CACHE, exist_ok=True)

# ---- tokens ----------------------------------------------------------
HEX = dict(black="000000", white="FFFFFF", text2="4C4A4A", rose50="F7F3F3", rose100="F4EBEB",
           rose200="E3D6D6", rose300="DED4D4", grey300="BFBFBF", magenta="EC008C", blue="31ABE2",
           orange="F9B804", red="E3012F", yellow="FFDC00", green="1BA660")
C = {k: RGBColor.from_string(v) for k, v in HEX.items()}
POLY = {
    "header": [(0, 0), (100, 0), (100, 55), (6, 100), (0, 82)],
    "plate": [(0, 3), (4, 0), (97, 2), (100, 6), (98, 94), (94, 100), (2, 97), (0, 93)],
    "card": [(0, 0), (100, 2), (98, 96), (95, 100), (3, 98), (0, 4)],
    "sub-card": [(2, 2), (99, 0), (100, 95), (97, 100), (0, 98)],
    "offer-card": [(0, 0), (97, 1), (100, 98), (2, 100)],
    "section-rose": [(0, 3), (62, 0), (100, 2), (100, 97), (38, 100), (0, 98)],
    "callout": [(0, 6), (5, 0), (92, 2), (100, 9), (97, 100), (3, 96)],
    "infoblock": [(0, 0), (100, 1), (99, 99), (0, 100)],
    "button": [(2, 5), (99, 0), (100, 88), (97, 100), (0, 95)],
    "case-image": [(0, 0), (97, 2), (100, 8), (99, 100), (0, 98)],
    "image-a": [(7, 0), (94, 2), (100, 10), (98, 90), (93, 100), (5, 97), (0, 88), (2, 8)],
    "image-b": [(6, 0), (97, 3), (100, 15), (96, 96), (88, 100), (2, 98), (0, 84), (4, 4)],
    "image-c": [(0, 5), (4, 0), (98, 2), (100, 12), (96, 100), (2, 96)],
}
HEAD, BODY, HEADM = "Roboto Condensed", "Roboto", "Roboto Condensed Medium"

SW, SH = 12192000, 6858000
L = 990000                 # samma vänsterkant som logotypen
R = SW - L
W = R - L
PX = 12700                 # 1 pt i EMU

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

# temats typsnitt -> Roboto Condensed / Roboto, så att ny text i PowerPoint blir rätt
theme = prs.slide_masters[0].part.part_related_by(RT.THEME)
theme._blob = theme.blob.replace(b'<a:latin typeface="Calibri Light"', b'<a:latin typeface="Roboto Condensed"') \
                        .replace(b'<a:latin typeface="Calibri"', b'<a:latin typeface="Roboto"')

# ---- tillgångar --------------------------------------------------------
def svg_png(rel, width=800):
    """SVG ur assets/ som PNG (PowerPoint via python-pptx tar inte SVG)."""
    import cairosvg
    out = os.path.join(CACHE, rel.replace("/", "_").replace(".svg", "_%d.png" % width))
    if not os.path.exists(out):
        cairosvg.svg2png(url=os.path.join(ASSETS, rel), write_to=out, output_width=width)
    return out


def bw(path):
    """Svartvitt foto: grayscale(1) contrast(1.05), max 2000 px."""
    out = os.path.join(CACHE, "bw_" + os.path.splitext(os.path.basename(path))[0] + ".jpg")
    if not os.path.exists(out):
        im = Image.open(path).convert("L")
        im.thumbnail((2000, 2000))
        ImageEnhance.Contrast(im).enhance(1.05).save(out, quality=86)
    return out


def P(name):
    """Foto ur assets/images/black-and-white/ (namn utan filändelse) eller sökväg till egen fil."""
    if os.path.isabs(name) or os.path.exists(name):
        return bw(name)
    folder = os.path.join(ASSETS, "images", "black-and-white")
    for ext in (".jpg", ".png", ".jpeg"):
        f = os.path.join(folder, name + ext)
        if os.path.exists(f):
            return bw(f)
    raise FileNotFoundError(name)


LOGO = svg_png("logos/magello-magenta.svg", 1200)
ICONLOGO = svg_png("logos/icon-logotype.svg", 600)
CITE = svg_png("graphics/citation.svg", 600)
ICON = {n: svg_png("graphics/%s.svg" % n, 400) for n in
        ["arrow-up-right", "cloud-arrow", "light", "markup", "nodes", "rocket", "speak-bouble", "thumbs-up", "timer"]}


# ---- primitiver -----------------------------------------------------------
def poly(s, x, y, w, h, pts, fill):
    pts = POLY[pts] if isinstance(pts, str) else pts
    fb = s.shapes.build_freeform(x + w * pts[0][0] / 100, y + h * pts[0][1] / 100, scale=1.0)
    fb.add_line_segments([(x + w * a / 100, y + h * b / 100) for a, b in pts[1:]], close=True)
    sh = fb.convert_to_shape()
    sh.fill.solid()
    sh.fill.fore_color.rgb = C[fill]
    sh.line.fill.background()
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)
    return sh


def poly_pic(s, path, x, y, w, h, pts="image-a"):
    pts = POLY[pts] if isinstance(pts, str) else pts
    x, y, w, h = int(x), int(y), int(w), int(h)
    pic = s.shapes.add_picture(path, x, y, w, h)
    iw, ih = Image.open(path).size
    t, src = w / h, iw / ih
    if src > t:
        k = t / src
        pic.crop_left = pic.crop_right = (1 - k) / 2
    else:
        k = src / t
        pic.crop_top = pic.crop_bottom = (1 - k) / 2
    spPr = pic._element.spPr
    cg = etree.Element(qn("a:custGeom"))
    for tag in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
        etree.SubElement(cg, qn(tag))
    etree.SubElement(cg, qn("a:rect"), l="0", t="0", r="r", b="b")
    pa = etree.SubElement(etree.SubElement(cg, qn("a:pathLst")), qn("a:path"), w="100000", h="100000")
    for i, (a, b) in enumerate(pts):
        el = etree.SubElement(pa, qn("a:moveTo" if i == 0 else "a:lnTo"))
        etree.SubElement(el, qn("a:pt"), x=str(int(a * 1000)), y=str(int(b * 1000)))
    etree.SubElement(pa, qn("a:close"))
    spPr.replace(spPr.find(qn("a:prstGeom")), cg)
    return pic


def img(s, path, x, y, h=None, w=None):
    iw, ih = Image.open(path).size
    if h is None:
        h = w * ih / iw
    if w is None:
        w = h * iw / ih
    return s.shapes.add_picture(path, int(x), int(y), int(w), int(h))


def icon(s, name, x, y, h):
    iw, ih = Image.open(ICON[name]).size
    w = h * iw / ih
    if w > h * 1.15:
        w = h * 1.15
        h = w * ih / iw
    return s.shapes.add_picture(ICON[name], int(x), int(y), int(w), int(h))


# textroller (pt): webbens skala, uppräknad för presentation
ROLE = {
    # Rubriker = DocTitle (Word) / h1 (webb): Roboto Condensed 500, svart. Aldrig magenta.
    "hero":    dict(font=HEADM, size=88, color="black", ls=0.92),
    "display": dict(font=HEADM, size=72, color="black", ls=0.95),
    "title":   dict(font=HEADM, size=60, color="black", ls=0.95),
    "h1":      dict(font=HEADM, size=48, color="black", ls=1.0),
    "sub":     dict(font=HEAD, size=26, color="text2", ls=1.3),
    "h3":      dict(font=HEAD, size=24, bold=True, color="black", ls=1.05),
    "body":    dict(font=BODY, size=20, color="black", ls=1.35),
    "lead":    dict(font=BODY, size=20, color="text2", ls=1.4),
    "small":   dict(font=BODY, size=18, color="black", ls=1.3),
    "kicker":  dict(font=HEAD, size=13, bold=True, color="magenta", caps=True, spc=182),
    "label":   dict(font=HEAD, size=13, bold=True, color="black", caps=True, spc=160),
    "crumb":   dict(font=BODY, size=13, color="grey300"),
    "quote":   dict(font=HEAD, size=44, bold=True, color="black", ls=1.0),
    "btn":     dict(font=HEADM, size=18, spc=72),
}


def runs_into(p, runs, role, **over):
    st = dict(ROLE[role]); st.update(over)
    p.alignment = st.get("align", PP_ALIGN.LEFT)
    p.line_spacing = st.get("ls", 1.0)
    for t in runs:
        txt, extra = (t, {}) if isinstance(t, str) else t
        r = p.add_run()
        r.text = txt
        f = r.font
        f.name = extra.get("font", st["font"])
        f.size = Pt(st["size"])
        f.bold = extra.get("bold", st.get("bold", False))
        f.color.rgb = C[extra.get("color", st["color"]) if "color" in extra or "color" in st else "black"]
        rPr = r._r.get_or_add_rPr()
        if st.get("caps"):
            rPr.set("cap", "all")
        if st.get("spc"):
            rPr.set("spc", str(st["spc"]))


def text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    """paras: list of (role, runs|str, {overrides}) ; before=pt space before"""
    tb = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, item in enumerate(paras):
        role, runs = item[0], item[1]
        over = item[2] if len(item) > 2 else {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(runs, str):
            runs = [runs]
        before = over.pop("before", 0)
        bullet = over.pop("bullet", False)
        runs_into(p, runs, role, **over)
        if before:
            p.space_before = Pt(before)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "228600"); pPr.set("indent", "-228600")
            bc = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:srgbClr")).set("val", HEX["black"])
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
            p.space_after = Pt(4)
    return tb


def h_of(role, lines):
    st = ROLE[role]
    return int(lines * st["size"] * st.get("ls", 1.0) * 1.2 * PX)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ---- webbkomponenter ---------------------------------------------------
def page(header=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C["white"]
    if header:
        poly(s, 0, 0, SW, 820000, "header", "rose100")
    return s


def rose_section(s, y0=1000000, y1=6600000):
    poly(s, 0, y0, SW, y1 - y0, "section-rose", "rose100")


def section_head(s, y, label, h2, p=None, w=7600000):
    text(s, L, y, W, 250000, [("label", label)])
    text(s, L, y + 330000, W, 600000, [("h2", h2)])
    if p:
        text(s, L, y + 950000, w, 900000, [("lead", p)])


def breadcrumb(s, y, parts):
    text(s, L, y, W, 220000, [("crumb", "   /   ".join(parts))])


def button(s, x, y, label, primary=True):
    w = int(len(label) * 18 * 0.52 * PX + 2 * 420000)
    h = 560000
    poly(s, x, y, w, h, "button", "magenta" if primary else "rose200")
    text(s, x, y, w, h, [("btn", label, dict(color="white" if primary else "black", align=PP_ALIGN.CENTER))],
         anchor=MSO_ANCHOR.MIDDLE)
    return x + w


def card(s, x, y, w, h, kind="card", fill="white", ic=None, h3=None, body=None, kicker=None, pad=300000,
         inline_icon=False, body_size=18):
    poly(s, x, y, w, h, kind, fill)
    cy = y + pad
    if ic and not inline_icon:
        icon(s, ic, x + pad, cy, 480000)
        cy += 650000
    paras = []
    if kicker:
        paras.append(("label", kicker, dict(color="text2")))
    if h3:
        paras.append(("h3", h3, dict(before=4 if kicker else 0)))
    if body:
        paras.append(("small", body, dict(before=6, size=body_size)))
    tx = x + pad
    if ic and inline_icon:
        icon(s, ic, x + pad, cy, 400000)
        tx = x + pad
        cy += 520000
    text(s, tx, cy, w - 2 * pad, y + h - cy - pad / 2, paras)


def row_card(s, x, y, w, h, ic, h3, body, fill="white", pad=280000, body_size=18):
    poly(s, x, y, w, h, "card", fill)
    icon(s, ic, x + pad, y + pad, 400000)
    tx = x + pad + 620000
    text(s, tx, y + pad - 20000, x + w - pad - tx, h - pad, [("h3", h3, dict(size=22)), ("small", body, dict(before=4, size=body_size))])


STEP = [("blue", "white"), ("orange", "black"), ("red", "white"), ("yellow", "black"), ("green", "white"), ("magenta", "white")]


def steps(s, x, y, w, h, labels, nums=None, size=16, gap=50000):
    n = len(labels)
    order = {2: [0, 5], 3: [0, 1, 5], 4: [0, 1, 4, 5], 6: [0, 1, 2, 3, 4, 5]}[n]
    cw = (w - gap * (n - 1)) / n
    tip = 230000
    for i, lab in enumerate(labels):
        fill, fg = STEP[order[i]]
        sx = x + i * (cw + gap)
        pts = [(1, 4), (100 - tip / cw * 100, 1), (100, 50), (100 - tip / cw * 100, 96), (2, 100)]
        if i:
            pts.append((tip * 0.8 / cw * 100, 50))
        poly(s, sx, y, cw, h, pts, fill)
        paras = []
        if nums:
            paras.append(("small", nums[i], dict(size=13, color=fg, font=HEAD, ls=1.0, align=PP_ALIGN.CENTER)))
        paras.append(("h3", lab, dict(size=size, color=fg, align=PP_ALIGN.CENTER, ls=1.0)))
        text(s, sx + (tip * 0.9 if i else 120000), y, cw - tip * (1.9 if i else 1.2), h, paras, anchor=MSO_ANCHOR.MIDDLE)


def infoblock(s, x, y, w, h, label, items, cols=3, size=18):
    poly(s, x, y, w, h, "infoblock", "rose50")
    pad = 380000
    text(s, x + pad, y + 280000, w - 2 * pad, 250000, [("kicker", label)])
    per = -(-len(items) // cols)
    cw = (w - 2 * pad - (cols - 1) * 300000) / cols
    for c in range(cols):
        chunk = items[c * per:(c + 1) * per]
        if chunk:
            text(s, x + pad + c * (cw + 300000), y + 640000, cw, h - 700000,
                 [("small", it, dict(size=size, bullet=True, ls=1.15)) for it in chunk])


def callout(s, x, y, w, h, body, label=None, size=22):
    poly(s, x, y, w, h, "callout", "rose200")
    paras = []
    if label:
        paras.append(("label", label))
    paras.append(("h3", body, dict(size=size, before=4 if label else 0)))
    text(s, x + 420000, y, w - 840000, h, paras, anchor=MSO_ANCHOR.MIDDLE)


def doc_title(s, crumb, title, sub, label):
    """Sidtitel (DocTitle) i rosagrå sektion: stor rubrik, underrubrik, liten etikett."""
    rose_section(s, 1050000, 6450000)
    breadcrumb(s, 1750000, crumb)
    text(s, L, 2150000, W, h_of("title", 2), [("title", title)])
    text(s, L, 4000000, 8200000, 900000, [("sub", sub)])
    text(s, L, 5150000, W, 250000, [("label", label, dict(color="grey300"))])


# =====================================================================
FULL = [(0, 0), (100, 0), (100, 100), (0, 100)]


def photo_page(photo):
    """Helbild i svartvitt. Lägg text på en vit platta (plate)."""
    s = page()
    poly_pic(s, P(photo), 0, 0, SW, SH, FULL)
    return s


def plate(s, x, y, w, h):
    poly(s, x - 450000, y, w + 900000, h, "plate", "white")


def item_grid(s, items, y, cols, colw, gap, rowh, body_size=18):
    """Ikon, rubrik och text utan bakgrund."""
    for i, (ic, h3, b) in enumerate(items):
        x = L + (i % cols) * (colw + gap)
        yy = y + (i // cols) * rowh
        icon(s, ic, x, yy, 480000)
        text(s, x, yy + 640000, colw, rowh - 700000,
             [("h3", h3), ("small", b, dict(before=6, size=body_size, color="text2"))])


def plain_list(s, x, y, w, items, cols=2, size=20, gap=500000):
    per = -(-len(items) // cols)
    cw = (w - (cols - 1) * gap) / cols
    for c in range(cols):
        chunk = items[c * per:(c + 1) * per]
        if chunk:
            text(s, x + c * (cw + gap), y, cw, per * 450000,
                 [("small", it, dict(size=size, bullet=True, ls=1.2)) for it in chunk])


def divider(photo, kicker, title, sub):
    s = photo_page(photo)
    pw, ph_ = 7600000, 3300000
    py = SH - ph_ - 650000
    plate(s, L, py, pw, ph_)
    text(s, L, py + 450000, pw, 300000, [("kicker", kicker)])
    text(s, L, py + 800000, pw, h_of("title", 2), [("title", title)])
    text(s, L, py + 800000 + h_of("title", 2) + 80000, pw, 600000, [("sub", sub, dict(size=24))])
    return s


# ---- ikonlogotypen uppe till höger, sida 2 och framåt -----------------------
IH = int(450000 * 1.7)                 # 70 % större än tidigare
IW = int(IH * 600 / 641)
IX, IY = SW - IW - 340000, 340000      # i övre högra hörnet, lika stort avstånd till topp och högerkant


def covered(slide):
    """Har platsen för ikonlogotypen en bakgrund (foto eller färgad yta)?"""
    x0, y0, x1, y1 = IX - 100000, IY - 100000, IX + IW + 100000, IY + IH + 100000
    for sh in slide.shapes:
        if sh.width is None or sh.width < 1500000:
            continue
        if sh.shape_type in (13, 5, 1) or sh.shape_type is None or True:
            if sh.left <= x0 and sh.top <= y0 and sh.left + sh.width >= x1 and sh.top + sh.height >= y1:
                if sh.has_text_frame and sh.shape_type == 17:
                    continue
                return True
    return False




def cover(photo, kicker, title, sub, date, nav=None, featured=None):
    """Förstasida: signaturytan (webbens header, uppförstorad) över ett svartvitt foto."""
    s = page()
    poly_pic(s, P(photo), 0, 0, SW, SH, FULL)
    kx = L / SW * 100
    poly(s, 0, 0, SW, 4650000, [(0, 0), (100, 0), (100, 55), (kx, 100), (0, 82)], "rose100")
    img(s, LOGO, L, 420000, h=420000)

    def tw(t, size=15):
        return int(len(t) * size * 0.5 * PX)

    nx = R
    if featured:
        fw = tw(featured)
        nx = R - fw - 50000
        text(s, nx, 470000, fw + 300000, 330000, [("small", featured, dict(size=15, font=HEADM, color="magenta", ls=1.0))])
        poly(s, nx, 800000, fw, 45000, [(0, 30), (100, 0), (100, 70), (0, 100)], "magenta")
    for t in reversed(nav or []):
        nx -= tw(t) + 500000
        text(s, nx, 470000, tw(t) + 300000, 330000, [("small", t, dict(size=15, font=HEADM, ls=1.0))])
    text(s, L, 1450000, 8000000, 300000, [("kicker", kicker)])
    text(s, L, 1720000, 9000000, h_of("hero", 1), [("hero", title, dict(size=96))])
    text(s, L, 3000000, 6000000, 500000, [("sub", sub)])
    text(s, L, 3650000, 3000000, 250000, [("label", date, dict(color="grey300"))])
    return s



def slide_up(slide):
    """Övergången "Push" uppåt (nästa bild glider upp nerifrån), som alla bilder ska ha."""
    sld = slide._element
    old = sld.find(qn("p:transition"))
    if old is not None:
        sld.remove(old)
    tr = etree.SubElement(sld, qn("p:transition"))
    tr.set("spd", "med")
    etree.SubElement(tr, qn("p:push")).set("dir", "u")
    cm = sld.find(qn("p:clrMapOvr"))
    anchor = cm if cm is not None else sld.find(qn("p:cSld"))
    anchor.addnext(tr)


def save(path, title=None):
    """Lägger ikonlogotypen uppe till höger från bild 2 (inte där platsen har foto eller färgad yta), sätter övergången på alla bilder och sparar."""
    for sl in list(prs.slides)[1:]:
        if not covered(sl):
            sl.shapes.add_picture(ICONLOGO, IX, IY, IW, IH)
    for sl in prs.slides:
        slide_up(sl)
    if title:
        prs.core_properties.title = title
    prs.save(path)
    return path

